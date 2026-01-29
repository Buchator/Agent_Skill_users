"""
Skills Platform - Servicio de Documentos
Maneja la carga y extracción de texto de documentos del usuario.
Formatos soportados: PDF, DOCX, TXT
"""
import os
import uuid
import shutil
from pathlib import Path
from typing import Dict, List, Optional
from fastapi import UploadFile

from config import settings


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


def _extract_text_from_pdf(file_path: Path) -> str:
    """Extrae texto de un archivo PDF"""
    from PyPDF2 import PdfReader
    reader = PdfReader(str(file_path))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n\n".join(text_parts)


def _extract_text_from_docx(file_path: Path) -> str:
    """Extrae texto de un archivo DOCX"""
    from docx import Document
    doc = Document(str(file_path))
    text_parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)
    # También extraer texto de tablas
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                text_parts.append(row_text)
    return "\n".join(text_parts)


def _extract_text_from_txt(file_path: Path) -> str:
    """Lee texto de un archivo TXT o MD"""
    return file_path.read_text(encoding="utf-8", errors="replace")


EXTRACTORS = {
    ".pdf": _extract_text_from_pdf,
    ".docx": _extract_text_from_docx,
    ".txt": _extract_text_from_txt,
    ".md": _extract_text_from_txt,
}


class DocumentService:
    """Servicio para manejar documentos cargados por el usuario"""

    def __init__(self):
        self.upload_dir = Path(settings.UPLOAD_DIRECTORY)
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    def _user_dir(self, user_id: int) -> Path:
        """Directorio de uploads para un usuario específico"""
        user_path = self.upload_dir / str(user_id)
        user_path.mkdir(parents=True, exist_ok=True)
        return user_path

    def validate_file(self, filename: str, size: int) -> Optional[str]:
        """Valida un archivo. Retorna mensaje de error o None si es válido."""
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED_EXTENSIONS:
            return f"Tipo de archivo no permitido: {ext}. Permitidos: {', '.join(ALLOWED_EXTENSIONS)}"
        if size > MAX_FILE_SIZE:
            max_mb = MAX_FILE_SIZE // (1024 * 1024)
            return f"El archivo excede el tamaño máximo de {max_mb} MB"
        return None

    async def save_upload(self, file: UploadFile, user_id: int) -> Dict:
        """
        Guarda un archivo subido y extrae su contenido de texto.
        Retorna dict con metadata y texto extraído.
        """
        filename = file.filename or "documento"
        ext = Path(filename).suffix.lower()

        # Leer contenido del archivo
        content = await file.read()

        # Validar
        error = self.validate_file(filename, len(content))
        if error:
            return {"success": False, "error": error}

        # Generar nombre único para evitar colisiones
        unique_name = f"{uuid.uuid4().hex}{ext}"
        user_path = self._user_dir(user_id)
        file_path = user_path / unique_name

        # Guardar archivo
        file_path.write_bytes(content)

        # Extraer texto
        try:
            extractor = EXTRACTORS.get(ext)
            if not extractor:
                file_path.unlink(missing_ok=True)
                return {"success": False, "error": f"No se puede procesar archivos {ext}"}

            extracted_text = extractor(file_path)

            if not extracted_text.strip():
                file_path.unlink(missing_ok=True)
                return {
                    "success": False,
                    "error": "No se pudo extraer texto del documento. Verifique que el archivo no esté vacío o protegido."
                }

            return {
                "success": True,
                "file_id": unique_name,
                "original_name": filename,
                "size": len(content),
                "text_content": extracted_text,
                "text_length": len(extracted_text),
            }
        except Exception as e:
            file_path.unlink(missing_ok=True)
            return {"success": False, "error": f"Error al procesar el documento: {str(e)}"}

    async def save_multiple_uploads(
        self, files: List[UploadFile], user_id: int
    ) -> Dict:
        """
        Procesa múltiples archivos subidos.
        Retorna dict con lista de resultados y texto combinado.
        """
        results = []
        all_texts = []

        for file in files:
            result = await self.save_upload(file, user_id)
            results.append(result)
            if result["success"]:
                all_texts.append(
                    f"=== Documento: {result['original_name']} ===\n{result['text_content']}"
                )

        combined_text = "\n\n".join(all_texts) if all_texts else ""
        successful = sum(1 for r in results if r["success"])

        return {
            "success": successful > 0,
            "total_files": len(files),
            "successful": successful,
            "failed": len(files) - successful,
            "results": results,
            "combined_text": combined_text,
        }

    def cleanup_user_files(self, user_id: int):
        """Limpia archivos temporales de un usuario"""
        user_path = self._user_dir(user_id)
        if user_path.exists():
            shutil.rmtree(user_path, ignore_errors=True)


# Instancia global
document_service = DocumentService()
