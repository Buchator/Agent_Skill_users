"""
Skills Platform - Servicio de Exportación a Word
Genera documentos DOCX formateados a partir de la respuesta de Claude.
"""
import io
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


class DocxExportService:
    """Genera documentos Word a partir de texto/markdown"""

    def generate_docx(self, content: str, title: str = "Documento generado") -> bytes:
        """
        Convierte contenido de texto/markdown a un documento DOCX formateado.
        Retorna los bytes del documento.
        """
        doc = Document()

        # Configurar estilos
        self._setup_styles(doc)

        # Agregar título del documento
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Agregar línea separadora
        doc.add_paragraph("_" * 60)

        # Parsear y agregar contenido
        self._parse_and_add_content(doc, content)

        # Agregar pie de página
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Documento generado por Skills Platform - DPH")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(150, 150, 150)

        # Guardar a bytes
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _setup_styles(self, doc: Document):
        """Configura estilos base del documento"""
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

    def _parse_and_add_content(self, doc: Document, content: str):
        """Parsea contenido markdown básico y lo agrega al documento"""
        lines = content.split("\n")
        i = 0
        in_table = False
        table_rows = []

        while i < len(lines):
            line = lines[i]

            # Detectar tablas markdown
            if "|" in line and not in_table:
                # Verificar si es una tabla (al menos la línea actual y la siguiente contienen |)
                if i + 1 < len(lines) and "|" in lines[i + 1]:
                    in_table = True
                    table_rows = []
                    table_rows.append(line)
                    i += 1
                    continue

            if in_table:
                if "|" in line:
                    # Saltar línea separadora de tabla (|---|---|)
                    if not re.match(r"^\|[\s\-:|]+\|$", line.strip()):
                        table_rows.append(line)
                    i += 1
                    continue
                else:
                    # Fin de la tabla, renderizarla
                    self._add_table(doc, table_rows)
                    in_table = False
                    table_rows = []
                    # No incrementar i, procesar esta línea normalmente
                    continue

            # Encabezados
            if line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)

            # Listas con viñetas
            elif re.match(r"^\s*[-*]\s+", line):
                text = re.sub(r"^\s*[-*]\s+", "", line)
                para = doc.add_paragraph(style="List Bullet")
                self._add_formatted_text(para, text)

            # Listas numeradas
            elif re.match(r"^\s*\d+\.\s+", line):
                text = re.sub(r"^\s*\d+\.\s+", "", line)
                para = doc.add_paragraph(style="List Number")
                self._add_formatted_text(para, text)

            # Línea horizontal
            elif re.match(r"^[-_*]{3,}$", line.strip()):
                doc.add_paragraph("_" * 60)

            # Línea vacía
            elif not line.strip():
                doc.add_paragraph()

            # Párrafo normal
            else:
                para = doc.add_paragraph()
                self._add_formatted_text(para, line)

            i += 1

        # Si terminamos dentro de una tabla, renderizarla
        if in_table and table_rows:
            self._add_table(doc, table_rows)

    def _add_formatted_text(self, paragraph, text: str):
        """Agrega texto con formato markdown básico (bold, italic) a un párrafo"""
        # Patrón para detectar **bold**, *italic*, y texto normal
        pattern = r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|([^*`]+))"
        matches = re.finditer(pattern, text)

        for match in matches:
            if match.group(2):  # **bold**
                run = paragraph.add_run(match.group(2))
                run.bold = True
            elif match.group(3):  # *italic*
                run = paragraph.add_run(match.group(3))
                run.italic = True
            elif match.group(4):  # `code`
                run = paragraph.add_run(match.group(4))
                run.font.name = "Consolas"
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(80, 80, 80)
            elif match.group(5):  # texto normal
                paragraph.add_run(match.group(5))

    def _add_table(self, doc: Document, rows: list):
        """Agrega una tabla al documento a partir de filas markdown"""
        if not rows:
            return

        # Parsear celdas
        parsed_rows = []
        for row in rows:
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        table.style = "Table Grid"

        for r_idx, row_data in enumerate(parsed_rows):
            for c_idx, cell_text in enumerate(row_data):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text.strip()
                    # Primera fila en bold (encabezado)
                    if r_idx == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True

        doc.add_paragraph()  # Espacio después de la tabla


# Instancia global
docx_export_service = DocxExportService()
